# -*- coding: utf-8 -*-
"""Dump specific cell ranges (values AND formulas) from a sheet for close inspection.
Uso: python dump_range.py <arquivo> <aba> <col_ini> <col_fim> <linha_ini> <linha_fim>
"""
import sys
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

def main():
    path, sheet, c1, c2, r1, r2 = sys.argv[1:7]
    r1, r2 = int(r1), int(r2)
    c1i, c2i = column_index_from_string(c1), column_index_from_string(c2)

    for data_only in (False, True):
        wb = openpyxl.load_workbook(path, data_only=data_only, read_only=True)
        ws = wb[sheet]
        tag = "FORMULA" if not data_only else "VALUE  "
        print(f"\n===== {sheet} [{tag}] linhas {r1}-{r2} colunas {c1}-{c2} =====")
        for r in range(r1, r2 + 1):
            rowvals = []
            for ci in range(c1i, c2i + 1):
                cell = ws.cell(row=r, column=ci)
                if cell.value is not None:
                    rowvals.append(f"{get_column_letter(ci)}{r}={cell.value!r}")
            if rowvals:
                print("  " + " | ".join(rowvals))
        wb.close()

if __name__ == "__main__":
    main()
