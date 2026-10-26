#!/usr/bin/env python3
"""Gera um .pptx a partir de um arquivo de dados JSON (sem texto de cliente embutido).

Uso: gerar_deck.py DADOS.json [SAIDA.pptx] [--perfil NOME] [--validar]
--validar só confere os dados (stdlib): estrutura e `[preencher` pendente; sai com 1 se houver problema.
Gerar o .pptx exige `pip install python-pptx` (dependência opcional, só deste adaptador).
"""
import json
import sys
from pathlib import Path

PERFIS = Path(__file__).with_name("perfis.json")


def problemas(dados, perfis):
    p = []
    if dados.get("perfil", "padrao") not in perfis:
        p.append(f"perfil '{dados.get('perfil')}' inexistente em perfis.json")
    textos = [dados.get("titulo", "")]
    if not dados.get("titulo"):
        p.append("titulo ausente")
    if not dados.get("slides"):
        p.append("slides ausente ou vazio")
    for i, s in enumerate(dados.get("slides", []), 1):
        if not s.get("titulo") or not (s.get("bullets") or s.get("tabela")):
            p.append(f"slide {i}: precisa de titulo e de bullets ou tabela")
        textos += [s.get("titulo", "")] + list(s.get("bullets", []))
        t = s.get("tabela") or {}
        textos += list(t.get("cabecalho", [])) + [c for l in t.get("linhas", []) for c in l]
    pend = sum("[preencher" in str(t) for t in textos)
    if pend:
        p.append(f"{pend} texto(s) com [preencher pendente")
    return p


def gerar(dados, perfil, saida):
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Pt
    except ImportError:
        sys.exit("python-pptx ausente: instale com `pip install python-pptx` para gerar o .pptx (ou use --validar)")
    cor = lambda h: RGBColor.from_string(h)
    prs = Presentation()
    prs.slide_width, prs.slide_height = 12192000, 6858000  # 16:9

    def estilo(slide, titulo_shape):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = cor(perfil["cor_fundo"])
        for r in titulo_shape.text_frame.paragraphs[0].runs:
            r.font.name, r.font.color.rgb = perfil["fonte"], cor(perfil["cor_titulo"])

    capa = prs.slides.add_slide(prs.slide_layouts[5])
    capa.shapes.title.text = dados["titulo"]
    estilo(capa, capa.shapes.title)
    for s in dados["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = s["titulo"]
        estilo(slide, slide.shapes.title)
        if s.get("bullets"):
            caixa = slide.shapes.add_textbox(prs.slide_width // 12, prs.slide_height // 4, prs.slide_width * 5 // 6, prs.slide_height // 2)
            for i, b in enumerate(s["bullets"]):
                par = caixa.text_frame.paragraphs[0] if i == 0 else caixa.text_frame.add_paragraph()
                run = par.add_run()
                run.text = "• " + b
                run.font.size, run.font.name, run.font.color.rgb = Pt(20), perfil["fonte"], cor(perfil["cor_texto"])
        else:
            t, cab = s["tabela"], s["tabela"]["cabecalho"]
            tab = slide.shapes.add_table(len(t["linhas"]) + 1, len(cab), prs.slide_width // 12, prs.slide_height // 4, prs.slide_width * 5 // 6, prs.slide_height // 2).table
            for j, texto in enumerate(cab):
                tab.cell(0, j).text = texto
            for i, linha in enumerate(t["linhas"], 1):
                for j, texto in enumerate(linha):
                    tab.cell(i, j).text = texto
    prs.save(saida)


def main(argv):
    args = [a for a in argv if not a.startswith("--")]
    perfil_arg = argv[argv.index("--perfil") + 1] if "--perfil" in argv else None
    if perfil_arg:
        args.remove(perfil_arg)
    if not args:
        print(__doc__.strip())
        return 2
    dados = json.loads(Path(args[0]).read_text(encoding="utf-8"))
    perfis = json.loads(PERFIS.read_text(encoding="utf-8"))
    if perfil_arg:
        dados["perfil"] = perfil_arg
    falhas = problemas(dados, perfis)
    for f in falhas:
        print(f"- {f}")
    if "--validar" in argv or falhas:
        print(f"{len(falhas)} problema(s)")
        return 1 if falhas else 0
    gerar(dados, perfis[dados.get("perfil", "padrao")], args[1] if len(args) > 1 else "deck.pptx")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
